from collections import Counter
from hashlib import sha256
from pathlib import Path
import sys
from openpyxl import load_workbook

if len(sys.argv) != 2:
    raise SystemExit('usage: audit_workbook.py PRIVATE_EXPORT.xlsx')
SOURCE = Path(sys.argv[1])
SPECS = {
    'Settings':'key', 'Members':'member_id', 'Sessions':'token',
    'MemberAddresses':'address_id', 'Products':'product_id',
    'Categories':'kategori_id', 'PickupLocations':'lokasi_id',
    'DeliverySlots':'slot_id', 'Holidays':'tanggal', 'Orders':'order_id',
    'OrderItems':None, 'ProductVariants':'variant_id',
    'ProductAddons':'addon_id', 'OrderItemAddons':'id',
    'PointHistory':'id', 'Reviews':'review_id', 'PromoCodes':'promo_id',
    'PromoUsage':'usage_id', 'MessageTemplates':'kode',
    'Campaigns':'campaign_id', 'Logs':None,
}
REFS = [
    ('MemberAddresses','member_id','Members','member_id'),
    ('Orders','member_id','Members','member_id'),
    ('OrderItems','order_id','Orders','order_id'),
    ('OrderItemAddons','order_id','Orders','order_id'),
    ('PointHistory','member_id','Members','member_id'),
    ('Reviews','order_id','Orders','order_id'),
    ('PromoUsage','order_id','Orders','order_id'),
    ('ProductVariants','product_id','Products','product_id'),
    ('ProductAddons','product_id','Products','product_id'),
]

wb = load_workbook(SOURCE, read_only=True, data_only=True)
print('business_tabs', len(SPECS), 'workbook_tabs', len(wb.sheetnames),
      'missing', sorted(set(SPECS)-set(wb.sheetnames)),
      'extra', sorted(set(wb.sheetnames)-set(SPECS)))
data = {}
for name, primary in SPECS.items():
    if name not in wb:
        continue
    rows = wb[name].iter_rows(values_only=True)
    headers = tuple(str(c).strip() if c is not None else '' for c in next(rows))
    active = [dict(zip(headers, r)) for r in rows if any(c is not None and str(c).strip() for c in r)]
    data[name] = active
    dup_headers = [k for k,n in Counter(h for h in headers if h).items() if n > 1]
    duplicate_ids = 0
    blank_ids = 0
    if primary:
        ids = [str(r.get(primary) or '').strip() for r in active]
        blank_ids = sum(not x for x in ids)
        duplicate_ids = sum(n-1 for n in Counter(x for x in ids if x).values() if n > 1)
    print(f'{name}: rows={len(active)} columns={len([x for x in headers if x])} '
          f'primary={primary or "none"} blank_ids={blank_ids} duplicate_ids={duplicate_ids} '
          f'duplicate_headers={len(dup_headers)}')

for child, ck, parent, pk in REFS:
    if child not in data or parent not in data:
        continue
    ids = {str(r.get(pk) or '').strip() for r in data[parent]}
    orphan = sum(bool(str(r.get(ck) or '').strip()) and str(r.get(ck) or '').strip() not in ids
                 for r in data[child])
    print(f'reference {child}.{ck}->{parent}.{pk}: orphan={orphan}')

addons = data.get('ProductAddons', [])
if addons:
    by_id = {}
    for sheet_row, row in enumerate(addons, 2):
        by_id.setdefault(str(row.get('addon_id') or '').strip(), []).append((sheet_row, row))
    for group in by_id.values():
        if len(group) > 1:
            fingerprints = {sha256(repr(sorted(row.items())).encode()).hexdigest() for _, row in group}
            print('ProductAddons duplicate group: sheet_rows=',
                  ','.join(str(n) for n,_ in group),
                  'identical_full_rows=', len(fingerprints)==1,
                  'distinct_products=', len({str(row.get('product_id')) for _,row in group}))

wb.close()
